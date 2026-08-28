import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

# Load template
prs = pptx.Presentation('CIT Hackthon.pptx')

# Colors
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = RGBColor(56, 189, 248)    # #38BDF8
MUTED_BLUE = RGBColor(148, 163, 184)   # #94A3B8
DARK_NAVY = RGBColor(15, 23, 42)      # #0F172A
CARD_BG = RGBColor(19, 28, 46)        # #131C2E
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399
ACCENT_PINK = RGBColor(244, 114, 182)  # #F472B6
ACCENT_YELLOW = RGBColor(251, 191, 36) # #FBBF24

def set_para_text(p, text, font_name="Segoe UI", font_size=20, color=WHITE, bold=False, italic=False, align=PP_ALIGN.LEFT):
    p.text = text
    p.alignment = align
    if p.runs:
        r = p.runs[0]
        r.font.name = font_name
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color

def add_bullet_point(tf, title, desc, title_color=LIGHT_BLUE, desc_color=WHITE, font_size=18, title_bold=True):
    p = tf.add_paragraph() if tf.paragraphs and tf.paragraphs[0].text else tf.paragraphs[0]
    p.space_after = Pt(10)
    p.line_spacing = 1.15
    
    r1 = p.add_run()
    r1.text = title + " "
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(font_size)
    r1.font.bold = title_bold
    r1.font.color.rgb = title_color
    
    r2 = p.add_run()
    r2.text = desc
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(font_size)
    r2.font.bold = False
    r2.font.color.rgb = desc_color

print("Script framework ready.")
